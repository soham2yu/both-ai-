"""
Answer-key / question-paper extraction.

Inputs: .docx, .pdf (typed or scanned) or .md/.txt. (Legacy binary .doc is
not supported — convert it to .docx first.)

Outputs are keyed by question number as a string ("1", "2", ...):

    answer key      {"1": {"answer": "...", "marks": 6.0}, ...}
    question paper  {"1": {"question": "...", "marks": 6.0 or None}, ...}

Marks live in the answer key (a deliberate design decision). If a mark value
can't be found it is left as None and reported — never guessed.

Parsing tiers, cheapest first:
  1. Regex on question markers + mark annotations (no network, instant)
  2. Ollama LLM, only if the regex result is empty or missing marks
"""

import re
import json

import pymupdf
import docx as python_docx

import llm_client
import pdf_ingest
import ocr as ocr_module

# Question markers. Groups: (1) "Q"/"Question" prefix present, (2) number after
# prefix, (3) bare number that must be followed by "." or ")".
# The leading class tolerates markdown decoration: "## Q1", "**Q1**", "> 1.".
QUESTION_MARKER = re.compile(
    r"(?:^|\n)[ \t#*>\-]*(?:(Q(?:uestion)?\.?)\s*(\d{1,3})|(\d{1,3})(?=[\.\)]))\s*[\.\):]?\s*",
    re.IGNORECASE,
)

# Mark annotations: "[6]", "(8 Marks)", "10M", "Marks: 10"
MARKS_PATTERN = re.compile(
    r"[\[\(]\s*(\d{1,3}(?:\.\d+)?)\s*(?:marks?|m)?\s*[\]\)]"
    r"|\b(\d{1,3}(?:\.\d+)?)\s*(?:marks?|m)\b"
    r"|\bmarks?\s*[:=]\s*(\d{1,3}(?:\.\d+)?)",
    re.IGNORECASE,
)
_MARKS_WINDOW = 60  # annotations sit at the start or end of a question, not mid-answer


# ---------------------------------------------------------------- raw text ---

def extract_raw_text(path: str) -> dict:
    """{"text", "source"}; source is docx | md | pdf_text_layer | pdf_ocr."""
    lower = path.lower()

    if lower.endswith((".md", ".markdown", ".txt")):
        with open(path, encoding="utf-8", errors="replace") as f:
            return {"text": f.read(), "source": "md"}

    if lower.endswith(".docx"):
        doc = python_docx.Document(path)
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        parts.append(cell.text.strip())
        return {"text": "\n".join(parts), "source": "docx"}

    if lower.endswith(".pdf"):
        doc = pymupdf.open(path)
        combined = "\n".join(page.get_text() for page in doc).strip()
        doc.close()
        if len(combined) > 20:
            return {"text": combined, "source": "pdf_text_layer"}
        pages = pdf_ingest.render_pdf_pages(path)  # scanned document: same OCR tiering as student sheets
        return {"text": "\n".join(ocr_module.ocr_image(p)["text"] for p in pages), "source": "pdf_ocr"}

    raise ValueError(f"Unsupported file type: {path} (expected .docx, .pdf, .md or .txt)")


# ------------------------------------------------------------ marker split ---

def _select_markers(matches: list) -> list:
    """Numbered lists INSIDE an answer look like question markers. If the
    document uses Q-prefixed markers, trust only those. Otherwise accept bare
    numbers only when they count up 1, 2, 3... in order."""
    prefixed = [m for m in matches if m.group(1)]
    if prefixed:
        return prefixed
    chosen, expected = [], None
    for m in matches:
        n = int(m.group(3))
        if expected is None or n == expected:
            chosen.append(m)
            expected = n + 1
    return chosen


def _split_by_markers(raw_text: str) -> list:
    """[(que_num, chunk_text), ...] in document order."""
    markers = _select_markers(list(QUESTION_MARKER.finditer(raw_text)))
    out = []
    for i, m in enumerate(markers):
        number = str(int(m.group(2) or m.group(3)))
        end = markers[i + 1].start() if i + 1 < len(markers) else len(raw_text)
        out.append((number, raw_text[m.end():end]))
    return out


def _pull_marks(chunk: str):
    """Returns (chunk_without_annotation, marks_or_None). Only looks near the
    start/end of the chunk so '3 m' inside an answer isn't misread as marks."""
    for lo, hi in ((0, _MARKS_WINDOW), (max(0, len(chunk) - _MARKS_WINDOW), len(chunk))):
        m = MARKS_PATTERN.search(chunk, lo, hi)
        if m:
            value = next(g for g in m.groups() if g is not None)
            return chunk[:m.start()] + chunk[m.end():], float(value)
    return chunk, None


def _clean(text: str) -> str:
    """Tidy whitespace but KEEP line breaks: numbered/bulleted lists inside an
    answer rely on them, and rubric_decompose splits list items on newlines."""
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\s*\n\s*", "\n", text)
    text = re.sub(r"^[\s\-\u2013\u2014:.;*#]+", "", text)   # leftovers like "- :" after a mark annotation is removed
    return text.strip(" .;:*#\n\t")


# ----------------------------------------------------------- tier 1 + tier 2 ---

def _regex_parse(raw_text: str) -> list:
    entries = []
    for number, chunk in _split_by_markers(raw_text):
        chunk, marks = _pull_marks(chunk)
        entries.append({"question_id": f"q{number}", "reference_answer": _clean(chunk), "total_marks": marks})
    return entries


def _llm_parse(raw_text: str) -> list:
    prompt = f"""The following is a faculty answer key for an exam. Extract every
question's number, its full reference answer text, and its mark value.

Return ONLY a JSON array, nothing else, in this exact shape:
[{{"question_id": "q1", "reference_answer": "...", "total_marks": 6}}, ...]

If a mark value truly isn't stated for a question, use null rather than guessing.

Document text:
\"\"\"
{raw_text[:12000]}
\"\"\"
"""
    parsed = llm_client.generate_json(prompt)
    if isinstance(parsed, list) and all(isinstance(e, dict) and "question_id" in e and "reference_answer" in e for e in parsed):
        return parsed
    return []


def parse_answer_key(path: str, use_llm_fallback: bool = True) -> dict:
    """List-shaped result used by the demo API:
    {"questions": [{"question_id","reference_answer","total_marks"}], "parse_method",
     "needs_manual_review", "source"}"""
    extracted = extract_raw_text(path)
    regex_result = _regex_parse(extracted["text"])
    incomplete = (not regex_result) or any(e["total_marks"] is None for e in regex_result)

    if incomplete and use_llm_fallback:
        llm_result = _llm_parse(extracted["text"])
        if llm_result:
            return {"questions": llm_result, "parse_method": "llm", "source": extracted["source"],
                    "needs_manual_review": any(e.get("total_marks") is None for e in llm_result)}

    return {"questions": regex_result, "parse_method": "regex_incomplete" if incomplete else "regex",
            "needs_manual_review": incomplete, "source": extracted["source"]}


# ------------------------------------------------------- structured outputs ---

def _qnum(question_id) -> str:
    digits = re.sub(r"\D", "", str(question_id))
    return str(int(digits)) if digits else str(question_id)


def build_answer_key(path: str, use_llm_fallback: bool = True) -> dict:
    """{"answer_key": {"1": {"answer", "marks"}}, "notes": [...], "parse_method", "source"}"""
    parsed = parse_answer_key(path, use_llm_fallback=use_llm_fallback)
    key, notes = {}, []
    for e in parsed["questions"]:
        n = _qnum(e["question_id"])
        marks = e.get("total_marks")
        key[n] = {"answer": e["reference_answer"], "marks": float(marks) if marks is not None else None}
        if not e["reference_answer"]:
            notes.append(f"Q{n}: answer text is empty")
        if marks is None:
            notes.append(f"Q{n}: marks not found — set 'marks' in answer_key.json before decomposing")
    if not key:
        notes.append("No questions found in the answer key document")
    return {"answer_key": key, "notes": notes, "parse_method": parsed["parse_method"], "source": parsed["source"]}


def build_question_paper(path: str) -> dict:
    """{"1": {"question": "...", "marks": 6.0 or None}, ...}"""
    raw = extract_raw_text(path)["text"]
    paper = {}
    for number, chunk in _split_by_markers(raw):
        chunk, marks = _pull_marks(chunk)
        paper[number] = {"question": _clean(chunk), "marks": marks}
    return paper


def cross_check(answer_key: dict, question_paper: dict) -> list:
    """Warnings where the question paper and answer key disagree."""
    notes = []
    for n in sorted(set(question_paper) - set(answer_key), key=int):
        notes.append(f"Q{n} is in the question paper but has no answer in the key")
    for n in sorted(set(answer_key) - set(question_paper), key=int):
        notes.append(f"Q{n} is in the answer key but not in the question paper")
    for n in sorted(set(answer_key) & set(question_paper), key=int):
        pm, km = question_paper[n].get("marks"), answer_key[n].get("marks")
        if pm is not None and km is not None and pm != km:
            notes.append(f"Q{n}: marks differ — paper says {pm}, answer key says {km} (key is used)")
    return notes


if __name__ == "__main__":
    import sys
    if len(sys.argv) > 1:
        print(json.dumps(build_answer_key(sys.argv[1]), indent=2))
    else:
        print("Usage: python3 answer_key_extractor.py <answer_key.pdf|docx|md>")
